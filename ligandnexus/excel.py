"""
excel.py — Exportación a Excel con estructuras 2D opcionales
============================================================
"""

from __future__ import annotations

import io
from rdkit import Chem
from rdkit.Chem import Draw


# CWE-1236. Excel interpreta como FORMULA toda celda que empiece por = + - @
# o por un caracter de control. Tanto los identificadores como los nombres de
# columna pueden venir de un archivo que sube el usuario, asi que un valor
# hostil se convierte en codigo que se ejecuta en la maquina de QUIEN ABRE el
# libro -- que normalmente no es quien lo genero, sino a quien se lo envia.
_INICIO_PELIGROSO = ("=", "+", "-", "@", "\t", "\r")


def _celda_segura(v):
    """Antepone un apostrofo para forzar que Excel lo trate como texto."""
    if isinstance(v, str) and v.startswith(_INICIO_PELIGROSO):
        return "'" + v
    return v


def exportar_excel(df, smiles_col, ruta_salida, titulo="derivados",
                   con_imagenes=True, img_size=(160, 120),
                   columnas=None, log=print):
    """Exporta `df` a Excel. `con_imagenes` renderiza la estructura 2D (lento)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.drawing.image import Image as XLImage

    if columnas is None:
        columnas = list(df.columns)

    IMG_W, IMG_H = img_size
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titulo[:31]

    HDR_FILL = PatternFill("solid", fgColor="15486B")
    HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
    LINK_F = Font(color="0563C1", underline="single", size=10)
    DATA_F = Font(size=10)
    CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

    col0 = 1
    headers = []
    if con_imagenes:
        headers.append("Estructura 2D")
        col0 = 2
    headers += columnas

    for ci, hdr in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=_celda_segura(hdr))
        c.fill = HDR_FILL; c.font = HDR_FONT; c.alignment = CTR
    ws.freeze_panes = "A2"

    if con_imagenes:
        ws.column_dimensions["A"].width = (IMG_W - 5) / 7 + 0.5

    total = len(df)
    for ri, (_, row) in enumerate(df.iterrows(), start=2):
        if con_imagenes:
            ws.row_dimensions[ri].height = IMG_H * 0.75
            smi = row.get(smiles_col, "")
            mol = Chem.MolFromSmiles(str(smi)) if smi else None
            if mol:
                img = Draw.MolToImage(mol, size=img_size, kekulize=True)
                buf = io.BytesIO(); img.save(buf, format="PNG"); buf.seek(0)
                ws.add_image(XLImage(buf), f"A{ri}")

        for c_off, col_name in enumerate(columnas, start=col0):
            val = row.get(col_name, "")
            cell = ws.cell(row=ri, column=c_off, value=_celda_segura(val))
            cell.font = DATA_F
            cell.alignment = LEFT if "smiles" in col_name.lower() or col_name == "IUPACName" else CTR
            # Solo enlazar a PubChem los IDs numéricos (los de ChEMBL/COCONUT no).
            if col_name == "CID" and str(val).isdigit():
                cell.hyperlink = f"https://pubchem.ncbi.nlm.nih.gov/compound/{val}"
                cell.font = LINK_F
        if (ri - 1) % 50 == 0:
            log(f"  Excel: {ri - 1}/{total} filas")

    ws.auto_filter.ref = ws.dimensions
    wb.save(ruta_salida)
    log(f"  Excel guardado.")
    return ruta_salida
